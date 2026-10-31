"""Tagesabschluss Lagerverkauf: Verkäufe aus der Kasse nach FEFO von den Losen abbuchen.

Die Kasse kennt nur Artikelnamen. Die Zuordnung Name → Artikel + Einheit steht in „Kassenartikel“
und wird beim Buchen für neue Namen gelernt. Abgebucht wird als Lagerbuchung „Abverkauf Lagerverkauf“
(Materialentnahme) aus dem Lagerort des Lagerverkaufs, je Artikel von den Losen mit dem frühesten MHD.
"""

import csv
import difflib
import io

import frappe
from frappe import _
from frappe.model.document import Document
from frappe.utils import flt, getdate, nowtime

from keks_trace import palette as paletten

BUCHUNGSART = "Abverkauf Lagerverkauf"
NAMEN_SPALTEN = ("artikel", "artikelname", "bezeichnung", "name", "produkt", "text")
MENGEN_SPALTEN = ("menge", "anzahl", "stück", "stueck", "stk", "qty")


class LagerverkaufAbschluss(Document):
	def validate(self):
		for p in self.positionen:
			if not p.artikel and p.kassenname:
				k = frappe.db.get_value(
					"Kassenartikel", {"kassenname": p.kassenname, "aktiv": 1}, ["artikel", "einheit"]
				)
				if k:
					p.artikel, p.einheit = k
			if p.artikel and not p.einheit:
				p.einheit = frappe.get_cached_value(
					"Item", p.artikel, "sales_uom"
				) or frappe.get_cached_value("Item", p.artikel, "stock_uom")

	def before_submit(self):
		offen = [p.kassenname or f"Zeile {p.idx}" for p in self.positionen if not p.artikel]
		if offen:
			frappe.throw(_("Bitte noch zuordnen: {0}").format(", ".join(offen)))
		self._zuordnungen_lernen()
		self.buchung = self._abbuchen()

	def on_cancel(self):
		if self.buchung:
			se = frappe.get_doc("Stock Entry", self.buchung)
			if se.docstatus == 1:
				se.cancel()

	def _zuordnungen_lernen(self):
		for p in self.positionen:
			if p.kassenname and not frappe.db.exists("Kassenartikel", p.kassenname):
				frappe.get_doc(
					{
						"doctype": "Kassenartikel",
						"kassenname": p.kassenname,
						"artikel": p.artikel,
						"einheit": p.einheit,
					}
				).insert(ignore_permissions=True)

	def _abbuchen(self) -> str:
		bedarf = {}
		for p in self.positionen:
			kg = flt(p.menge) * paletten.faktor(p.artikel, p.einheit)
			bedarf[p.artikel] = bedarf.get(p.artikel, 0) + kg
		zeilen, fehlt = [], []
		for artikel, menge in bedarf.items():
			rest = menge
			for charge, verfuegbar in _lose_fefo(artikel, self.lagerort):
				nimm = min(rest, verfuegbar)
				if nimm <= 1e-6:
					continue
				uom = frappe.get_cached_value("Item", artikel, "stock_uom")
				zeilen.append(
					{
						"item_code": artikel,
						"qty": nimm,
						"transfer_qty": nimm,
						"uom": uom,
						"stock_uom": uom,
						"conversion_factor": 1,
						"s_warehouse": self.lagerort,
						"use_serial_batch_fields": 1,
						"batch_no": charge,
					}
				)
				for p in self.positionen:
					if p.artikel == artikel:
						p.chargen = ((p.chargen or "") + f"{charge}: {flt(nimm, 3):g}\n").strip()
				rest -= nimm
				if rest <= 1e-6:
					break
			if rest > 1e-4:
				fehlt.append(
					_("{0}: {1} fehlen im {2}").format(
						artikel, paletten.zerlegen(artikel, rest), self.lagerort
					)
				)
		if fehlt:
			frappe.throw(
				_("Im Lagerverkauf ist nicht genug Ware gebucht (vorher umlagern?): {0}").format(
					"; ".join(fehlt)
				)
			)
		se = frappe.get_doc(
			{
				"doctype": "Stock Entry",
				"stock_entry_type": BUCHUNGSART,
				"purpose": "Material Issue",
				"company": self.company,
				"set_posting_time": 1,
				"posting_date": getdate(self.datum),
				"posting_time": nowtime() if getdate(self.datum) == getdate() else "20:00:00",
				"remarks": _("Lagerverkauf {0}").format(self.name),
				"items": zeilen,
			}
		)
		se.insert(ignore_permissions=True)
		se.submit()
		return se.name


def _lose_fefo(artikel, lagerort):
	"""Lose eines Artikels im Lagerort, frühestes MHD zuerst."""
	return [
		(z.charge, flt(z.menge))
		for z in frappe.db.sql(
			"""
			select sbe.batch_no as charge, sum(sbe.qty) as menge, b.expiry_date
			from `tabSerial and Batch Entry` sbe
			join `tabSerial and Batch Bundle` sbb on sbb.name = sbe.parent
			join `tabBatch` b on b.name = sbe.batch_no
			where sbe.warehouse = %(lager)s and b.item = %(artikel)s and sbb.is_cancelled = 0 and sbb.docstatus = 1
				and ifnull(b.kt_status, '') != 'Gesperrt'
			group by sbe.batch_no, b.expiry_date
			having sum(sbe.qty) > 0.000001
			order by b.expiry_date is null, b.expiry_date, sbe.batch_no
			""",
			{"lager": lagerort, "artikel": artikel},
			as_dict=True,
		)
	]


@frappe.whitelist()
def einlesen(name: str) -> dict:
	"""Export der Kasse einlesen: Positionen aus Name + Anzahl, bekannte Namen direkt zuordnen."""
	doc = frappe.get_doc("Lagerverkauf Abschluss", name)
	doc.check_permission("write")
	if doc.docstatus != 0:
		frappe.throw(_("Nur im Entwurf möglich."))
	if not doc.datei:
		frappe.throw(_("Bitte zuerst den Export der Kasse anhängen."))
	zeilen = _tabelle(frappe.get_doc("File", {"file_url": doc.datei}))
	summen = {}
	for name_, menge in zeilen:
		summen[name_] = summen.get(name_, 0) + menge
	bekannte = {
		k.kassenname: k
		for k in frappe.get_all("Kassenartikel", {"aktiv": 1}, ["kassenname", "artikel", "einheit"])
	}
	artikelnamen = {
		i.item_name: i.name
		for i in frappe.get_all(
			"Item", {"kt_chargen_stufe": "Fertigware", "disabled": 0}, ["name", "item_name"]
		)
	}
	doc.set("positionen", [])
	unbekannt = 0
	for name_, menge in summen.items():
		k = bekannte.get(name_)
		vorschlag = None
		if not k:
			unbekannt += 1
			treffer = difflib.get_close_matches(name_, list(bekannte) + list(artikelnamen), n=1, cutoff=0.45)
			if treffer:
				t = treffer[0]
				vorschlag = f"{t} → {bekannte[t].artikel}" if t in bekannte else f"{t} ({artikelnamen[t]})"
		doc.append(
			"positionen",
			{
				"kassenname": name_,
				"artikel": k.artikel if k else None,
				"einheit": k.einheit if k else None,
				"menge": menge,
				"vorschlag": vorschlag,
			},
		)
	doc.save()
	return {"positionen": len(summen), "unbekannt": unbekannt}


def _tabelle(datei) -> list[tuple[str, float]]:
	inhalt = datei.get_content()
	if datei.file_name.lower().endswith((".xlsx", ".xlsm")):
		from openpyxl import load_workbook

		ws = load_workbook(io.BytesIO(inhalt), read_only=True, data_only=True).active
		reihen = [[("" if c is None else str(c)).strip() for c in r] for r in ws.iter_rows(values_only=True)]
	else:
		text = inhalt.decode("utf-8-sig", errors="replace") if isinstance(inhalt, bytes) else inhalt
		try:
			dialekt = csv.Sniffer().sniff(text[:2048], delimiters=";,\t")
		except csv.Error:
			dialekt = csv.excel
			dialekt.delimiter = ";"
		reihen = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialekt)]
	reihen = [r for r in reihen if any(r)]
	if not reihen:
		frappe.throw(_("Die Datei ist leer."))
	kopf = [c.lower() for c in reihen[0]]
	sp_name = next((i for i, c in enumerate(kopf) if c in NAMEN_SPALTEN), None)
	sp_menge = next((i for i, c in enumerate(kopf) if c in MENGEN_SPALTEN), None)
	if sp_name is None or sp_menge is None:
		frappe.throw(
			_(
				"Spalten nicht erkannt. Erwartet eine Spalte für den Namen ({0}) und eine für die Anzahl ({1})."
			).format(", ".join(NAMEN_SPALTEN), ", ".join(MENGEN_SPALTEN))
		)
	ergebnis = []
	for r in reihen[1:]:
		if len(r) <= max(sp_name, sp_menge) or not r[sp_name]:
			continue
		try:
			menge = (
				float(r[sp_menge].replace(".", "").replace(",", "."))
				if "," in r[sp_menge]
				else float(r[sp_menge])
			)
		except ValueError:
			continue
		if menge:
			ergebnis.append((r[sp_name], menge))
	return ergebnis
